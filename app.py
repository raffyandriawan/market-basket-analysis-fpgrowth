# app.py

from flask import Flask, render_template, request, jsonify, send_file
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import func, distinct
from sqlalchemy import text   
from config import Config
from models import db, Transaksi, UploadLog
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import pandas as pd
import os
from datetime import datetime
from mlxtend.frequent_patterns import fpgrowth, association_rules
from mlxtend.preprocessing import TransactionEncoder
import io
import csv
import io
    

app = Flask(__name__)
app.config.from_object(Config)

db.init_app(app)


os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)


# ============================================================
# HALAMAN MANAGE DATA
# ============================================================

@app.route('/')
@app.route('/manage-data')
def manage_data():
    return render_template('manage_data.html')


@app.route('/api/transaksi', methods=['GET'])
def get_transaksi():
    """Ambil data transaksi dengan pagination dan search."""
    page     = request.args.get('page', 1, type=int)
    per_page = request.args.get('per_page', 10, type=int)
    search   = request.args.get('search', '', type=str)

    query = Transaksi.query
    if search:
        query = query.filter(
            db.or_(
                Transaksi.no_invoice.like(f'%{search}%'),
                Transaksi.nama_produk.like(f'%{search}%'),
                Transaksi.kode_produk.like(f'%{search}%')
            )
        )

    total    = query.count()
    transaksi = query.order_by(Transaksi.tanggal_transaksi.asc(), Transaksi.no_invoice.asc()) \
                     .paginate(page=page, per_page=per_page, error_out=False)

    return jsonify({
        'data'       : [t.to_dict() for t in transaksi.items],
        'total'      : total,
        'page'       : page,
        'per_page'   : per_page,
        'total_pages': transaksi.pages
    })


@app.route('/api/upload', methods=['POST'])
def upload_file():
    """Upload file Excel dan simpan ke database."""
    if 'file' not in request.files:
        return jsonify({'success': False, 'message': 'Tidak ada file yang dipilih'}), 400

    file = request.files['file']
    if file.filename == '':
        return jsonify({'success': False, 'message': 'Tidak ada file yang dipilih'}), 400

    if not file.filename.endswith(('.xlsx', '.xls')):
        return jsonify({'success': False, 'message': 'Format file harus Excel (.xlsx / .xls)'}), 400

    try:
        # Baca Excel
        df = pd.read_excel(file, dtype={
            'no_invoice'  : str,
            'kode_produk' : str,
            'nama_produk' : str,
        })

        # Normalisasi nama kolom (lowercase, strip spasi)
        df.columns = [c.strip().lower() for c in df.columns]

        # Validasi kolom
        required_cols = ['no_invoice', 'kode_produk', 'nama_produk', 'tanggal_transaksi']
        missing = [c for c in required_cols if c not in df.columns]
        if missing:
            return jsonify({
                'success': False,
                'message': f'Kolom tidak ditemukan: {", ".join(missing)}'
            }), 400

        # Bersihkan data
        df = df[required_cols].dropna()
        df['no_invoice']   = df['no_invoice'].str.strip()
        df['kode_produk']  = df['kode_produk'].str.strip()
        df['nama_produk']  = df['nama_produk'].str.strip()

        # Konversi tanggal
        df['tanggal_transaksi'] = pd.to_datetime(
            df['tanggal_transaksi'], dayfirst=True, errors='coerce'
        )
        df = df.dropna(subset=['tanggal_transaksi'])

        # Simpan ke database
        rows = []
        for _, row in df.iterrows():
            rows.append(Transaksi(
                no_invoice        = row['no_invoice'],
                kode_produk       = row['kode_produk'],
                nama_produk       = row['nama_produk'],
                tanggal_transaksi = row['tanggal_transaksi'].date()
            ))

        db.session.bulk_save_objects(rows)

        # Simpan log upload
        log = UploadLog(
            nama_file    = file.filename,
            jumlah_baris = len(rows)
        )
        db.session.add(log)
        db.session.commit()

        return jsonify({
            'success': True,
            'message': f'Berhasil mengupload {len(rows)} baris data dari {file.filename}'
        })

    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'Error: {str(e)}'}), 500


@app.route('/api/transaksi/<int:id>', methods=['DELETE'])
def delete_transaksi(id):
    """Hapus satu baris transaksi."""
    t = Transaksi.query.get_or_404(id)
    db.session.delete(t)
    db.session.commit()
    return jsonify({'success': True, 'message': 'Data berhasil dihapus'})


@app.route('/api/transaksi/delete-all', methods=['DELETE'])
def delete_all_transaksi():
    """Hapus semua data transaksi."""
    Transaksi.query.delete()
    db.session.commit()
    return jsonify({'success': True, 'message': 'Semua data berhasil dihapus'})


@app.route('/api/transaksi', methods=['POST'])
def insert_transaksi():
    """Insert satu baris transaksi secara manual."""
    data = request.get_json()
    try:
        t = Transaksi(
            no_invoice        = data['no_invoice'].strip(),
            kode_produk       = data['kode_produk'].strip(),
            nama_produk       = data['nama_produk'].strip(),
            tanggal_transaksi = datetime.strptime(data['tanggal_transaksi'], '%Y-%m-%d').date()
        )
        db.session.add(t)
        db.session.commit()
        return jsonify({'success': True, 'message': 'Data berhasil ditambahkan'})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'Error: {str(e)}'}), 500


@app.route('/api/transaksi/<int:id>', methods=['PUT'])
def edit_transaksi(id):
    """Edit satu baris transaksi."""
    t    = Transaksi.query.get_or_404(id)
    data = request.get_json()
    try:
        t.no_invoice        = data['no_invoice'].strip()
        t.kode_produk       = data['kode_produk'].strip()
        t.nama_produk       = data['nama_produk'].strip()
        t.tanggal_transaksi = datetime.strptime(data['tanggal_transaksi'], '%Y-%m-%d').date()
        db.session.commit()
        return jsonify({'success': True, 'message': 'Data berhasil diupdate'})
    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'message': f'Error: {str(e)}'}), 500


@app.route('/api/transaksi/<int:id>', methods=['GET'])
def get_transaksi_by_id(id):
    """Ambil satu transaksi by ID (untuk form edit)."""
    t = Transaksi.query.get_or_404(id)
    return jsonify({'success': True, 'data': t.to_dict()})


@app.route('/api/upload-log', methods=['GET'])
def get_upload_log():
    """Ambil riwayat upload file."""
    logs = UploadLog.query.order_by(UploadLog.uploaded_at.desc()).all()
    return jsonify({'data': [l.to_dict() for l in logs]})


@app.route('/api/stats', methods=['GET'])
def get_stats():
    """Ambil statistik ringkasan data transaksi."""
    total_invoice = db.session.query(
    func.count(
        distinct(
            func.concat(Transaksi.no_invoice, '_', Transaksi.tanggal_transaksi)
             )
        )
    ).scalar()
    total_produk  = db.session.query(db.func.count(db.distinct(Transaksi.nama_produk))).scalar()
    min_date      = db.session.query(db.func.min(Transaksi.tanggal_transaksi)).scalar()
    max_date      = db.session.query(db.func.max(Transaksi.tanggal_transaksi)).scalar()

    periode = '-'
    if min_date and max_date:
        periode = f"{min_date.strftime('%Y')}–{max_date.strftime('%Y')}"

    return jsonify({
        'total_invoice': total_invoice or 0,
        'total_produk' : total_produk or 0,
        'periode'      : periode
    })


# ============================================================
# HALAMAN VIEW RESULT
# ============================================================

@app.route('/view-result')
def view_result():
    # Ambil rentang tanggal dari database untuk default filter
    min_date = db.session.query(db.func.min(Transaksi.tanggal_transaksi)).scalar()
    max_date = db.session.query(db.func.max(Transaksi.tanggal_transaksi)).scalar()
    return render_template('view_result.html',
        min_date=min_date.strftime('%Y-%m-%d') if min_date else '',
        max_date=max_date.strftime('%Y-%m-%d') if max_date else ''
    )


@app.route('/api/run-fpgrowth', methods=['POST'])
def run_fpgrowth():
    """Jalankan algoritma FP-Growth."""
    data           = request.get_json()
    min_support    = float(data.get('min_support', 0.05))
    min_confidence = float(data.get('min_confidence', 0.5))
    date_start     = data.get('date_start', '')
    date_end       = data.get('date_end', '')

    try:
        # Query data sesuai filter tanggal
        query = Transaksi.query
        if date_start:
            query = query.filter(Transaksi.tanggal_transaksi >= datetime.strptime(date_start, '%Y-%m-%d').date())
        if date_end:
            query = query.filter(Transaksi.tanggal_transaksi <= datetime.strptime(date_end, '%Y-%m-%d').date())

        transaksi = query.all()

        if not transaksi:
            return jsonify({'success': False, 'message': 'Tidak ada data pada rentang waktu tersebut'}), 400

        # Bentuk itemset: group by no_invoice → list nama_produk
        df = pd.DataFrame([t.to_dict() for t in transaksi])
        df['invoice_key'] = df['no_invoice'] + '_' + df['tanggal_transaksi'].astype(str)
        basket = df.groupby('invoice_key')['nama_produk'].apply(list).tolist()

        # Filter transaksi yang hanya punya 1 item (tidak bisa membentuk asosiasi)
        basket = [list(set(items)) for items in basket if len(set(items)) > 1]

        if not basket:
            return jsonify({'success': False, 'message': 'Tidak ada transaksi dengan lebih dari 1 produk berbeda'}), 400

        # TransactionEncoder
        te      = TransactionEncoder()
        te_array = te.fit(basket).transform(basket)
        df_enc  = pd.DataFrame(te_array, columns=te.columns_)

        # FP-Growth
        frequent_itemsets = fpgrowth(df_enc, min_support=min_support, use_colnames=True)

        if frequent_itemsets.empty:
            return jsonify({'success': False, 'message': 'Tidak ada frequent itemset ditemukan. Coba turunkan nilai minimum support'}), 400

        # Association Rules
        rules = association_rules(frequent_itemsets, metric='confidence', min_threshold=min_confidence)

        if rules.empty:
            return jsonify({'success': False, 'message': 'Tidak ada aturan asosiasi ditemukan. Coba turunkan nilai minimum confidence'}), 400

        # Format hasil
        rules = rules.sort_values('confidence', ascending=False)
        rules['antecedents'] = rules['antecedents'].apply(lambda x: ', '.join(sorted(x)))
        rules['consequents'] = rules['consequents'].apply(lambda x: ', '.join(sorted(x)))
        rules['support']     = rules['support'].round(4)
        rules['confidence']  = rules['confidence'].round(4)
        rules['lift']        = rules['lift'].round(4)

        result = rules[['antecedents', 'consequents', 'support', 'confidence', 'lift']].to_dict('records')

        # Rekomendasi bundling: ambil top 5 berdasarkan confidence tertinggi
        rekomendasi = []
        for r in result[:5]:
            rekomendasi.append({
                'produk'    : f"{r['antecedents']} + {r['consequents']}",
                'confidence': f"{r['confidence']*100:.1f}%",
                'support'   : f"{r['support']*100:.1f}%",
                'lift'      : r['lift']
            })

        return jsonify({
            'success'    : True,
            'total_rules': len(result),
            'total_transaksi': len(basket),
            'rules'      : result,
            'rekomendasi': rekomendasi
        })

    except Exception as e:
        return jsonify({'success': False, 'message': f'Error: {str(e)}'}), 500


@app.route('/api/export-csv', methods=['POST'])
def export_csv():
    """Export hasil FP-Growth ke CSV."""
    data  = request.get_json()
    rules = data.get('rules', [])

    if not rules:
        return jsonify({'success': False, 'message': 'Tidak ada data untuk diexport'}), 400

    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=['antecedents', 'consequents', 'support', 'confidence', 'lift'])
    writer.writeheader()
    writer.writerows(rules)

    output.seek(0)
    return send_file(
        io.BytesIO(output.getvalue().encode('utf-8')),
        mimetype='text/csv',
        as_attachment=True,
        download_name=f'fp_growth_result_{datetime.now().strftime("%Y%m%d_%H%M%S")}.csv'
    )

@app.route('/api/download-template', methods=['GET'])
def download_template():
    """Download template Excel dengan data dummy."""

    wb = Workbook()
    ws = wb.active
    ws.title = 'Template Transaksi'

    # Header
    headers = ['no_invoice', 'kode_produk', 'nama_produk', 'tanggal_transaksi']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', start_color='0000CD')
        cell.alignment = Alignment(horizontal='center')

    # Data dummy
    dummy_data = [
        ['1000001', '179 400 000', 'BIBIT AYAM (NDK+NDL+VIT+IBD)', '01/01/2019'],
        ['1000002', '179 400 000', 'BIBIT AYAM (NDK+NDL+VIT+IBD)', '01/01/2019'],
        ['1000003', '179 400 000', 'BIBIT AYAM (NDK+NDL+VIT+IBD)', '01/01/2019'],
        ['1000004', '179 400 999', 'BIBIT AYAM PREMIUM',           '01/01/2019'],
        ['1000005', '179 400 000', 'BIBIT AYAM (NDK+NDL+VIT+IBD)', '02/01/2019'],
    ]

    for row_idx, row in enumerate(dummy_data, 2):
        for col_idx, value in enumerate(row, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    # Lebar kolom
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 35
    ws.column_dimensions['D'].width = 20

    # Simpan ke memory
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='template_transaksi.xlsx'
    )
@app.route('/api/produk', methods=['GET'])
def get_produk():
    """Ambil semua produk untuk dropdown."""
    from sqlalchemy import text
    hasil = db.session.execute(text('SELECT kode_produk, nama_produk FROM produk ORDER BY nama_produk')).fetchall()
    return jsonify({'data': [{'kode_produk': r[0], 'nama_produk': r[1]} for r in hasil]})

if __name__ == '__main__':
    with app.app_context():
        db.create_all()
    app.run(debug=True)